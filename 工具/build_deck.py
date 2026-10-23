# -*- coding: utf-8 -*-
"""建立「磁鐵好好玩」教學簡報（24 頁，16:9，米白風格）"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image

ROOT = r'C:\Users\rayi\Documents\自然科學教學專欄\磁鐵好好玩'
IMG = os.path.join(ROOT, '教學網頁', 'media', 'img')
OUT = os.path.join(ROOT, '簡報', '磁鐵好好玩_教學簡報.pptx')

BG = RGBColor(0xFA, 0xF7, 0xF0)
INK = RGBColor(0x35, 0x54, 0x4A)
INK_SOFT = RGBColor(0x6B, 0x7B, 0x72)
ORANGE = RGBColor(0xC0, 0x6A, 0x2B)
NRED = RGBColor(0xC0, 0x39, 0x2B)
SBLUE = RGBColor(0x2C, 0x5F, 0x8A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = '微軟正黑體'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = 13.333, 7.5


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background(); bg.shadow.inherit = False
    return s


def tb(s, x, y, w, h, text, size=24, color=INK, bold=False, align=PP_ALIGN.LEFT,
       spacing=1.35, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    lines = text.split('\n')
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = FONT
        rPr = r.font._rPr
        rPr.set('lang', 'zh-TW')
        # 一定要補 a:ea，否則 LibreOffice／PDF 會用替代字型，匯出中文會壞字
        latin = rPr.find(qn('a:latin'))
        ea = rPr.find(qn('a:ea'))
        if ea is None:
            ea = OxmlElement('a:ea')
            (latin.addnext(ea) if latin is not None else rPr.append(ea))
        ea.set('typeface', FONT)
        cs = rPr.find(qn('a:cs'))
        if cs is None:
            cs = OxmlElement('a:cs')
            (ea.addnext(cs) if ea is not None else rPr.append(cs))
        cs.set('typeface', FONT)
    return box


def rule(s, y, w=2.2, color=ORANGE, x=0.9):
    bar = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Pt(4))
    bar.fill.solid(); bar.fill.fore_color.rgb = color
    bar.line.fill.background(); bar.shadow.inherit = False


def pic(s, code, x, y, w, h, border=True):
    p = os.path.join(IMG, 'm%s.jpg' % code)
    if not os.path.exists(p):
        return
    im = Image.open(p); iw, ih = im.size
    sc = min(w / iw, h / ih)
    tw, th = iw * sc, ih * sc
    shape = s.shapes.add_picture(p, Inches(x + (w - tw) / 2), Inches(y + (h - th) / 2),
                                 Inches(tw), Inches(th))
    if border:
        shape.line.color.rgb = RGBColor(0xE2, 0xDA, 0xC8)
        shape.line.width = Pt(1)


def pic_grid(s, codes, x, y, w, h, cols=None, gap=0.22, label=None):
    n = len(codes)
    cols = cols or min(n, 4)
    rows = (n + cols - 1) // cols
    cw = (w - gap * (cols - 1)) / cols
    ch = (h - gap * (rows - 1)) / rows
    for i, c in enumerate(codes):
        r, col = divmod(i, cols)
        pic(s, c, x + col * (cw + gap), y + r * (ch + gap), cw, ch)
    if label:
        tb(s, x, y + h + 0.12, w, 0.4, label, 14, INK_SOFT, align=PP_ALIGN.CENTER)


def head(s, no, title, sub=None):
    tb(s, 0.9, 0.62, 1.2, 0.7, no, 20, ORANGE, True)
    tb(s, 1.55, 0.5, 10.9, 0.9, title, 34, INK, True)
    rule(s, 1.42)
    if sub:
        tb(s, 1.55, 1.6, 10.9, 0.6, sub, 18, INK_SOFT)


def bullets(s, items, x=1.0, y=1.9, w=10.8, size=24, gap=1.9):
    for i, t in enumerate(items):
        dot = s.shapes.add_shape(9, Inches(x), Inches(y + i * gap + 0.09), Inches(0.13), Inches(0.13))
        dot.fill.solid(); dot.fill.fore_color.rgb = ORANGE
        dot.line.fill.background(); dot.shadow.inherit = False
        tb(s, x + 0.35, y + i * gap, w, 0.7, t, size, INK)


def footer(s, n):
    tb(s, 0.9, 6.95, 8, 0.4, '磁鐵好好玩｜翰林三上自然 第2單元', 11, RGBColor(0xA9, 0xA0, 0x8E))
    tb(s, 11.4, 6.95, 1.0, 0.4, str(n), 11, RGBColor(0xA9, 0xA0, 0x8E), align=PP_ALIGN.RIGHT)


# ── 1 封面 ──────────────────────────────────────────────
s = slide()
tb(s, 0.9, 1.5, 6, 0.5, '翰林版｜三年級上學期｜自然科學', 18, INK_SOFT)
tb(s, 0.9, 2.05, 7.2, 1.5, '磁鐵好好玩', 66, INK, True)
rule(s, 3.55, 2.6)
tb(s, 0.9, 3.85, 7, 0.6, '第 2 單元｜4 節課｜磁力・磁極・生活應用', 24, INK_SOFT)
pic(s, '001', 8.4, 1.6, 3.6, 3.4)
tb(s, 0.9, 6.3, 8, 0.5, '115 學年度　10/26、10/29、11/2、11/5', 16, RGBColor(0xA9, 0xA0, 0x8E))

# ── 2 單元地圖 ──────────────────────────────────────────
s = slide(); head(s, 'MAP', '這 4 節課，我們要一起做這些事')
rows = [('1', '10/26 一', '磁鐵能吸引什麼？', '預測 → 實驗 → 歸納「鐵製品」'),
        ('2', '10/29 四', '磁力有多大？磁極在哪裡？', '兩端最強、越遠越弱、找出 N／S'),
        ('3', '11/2 一', '同極相斥、異極相吸', '相斥相吸、磁鐵會指南北'),
        ('4', '11/5 四', '生活中的磁鐵', '應用辨識、測驗 10 題、線上對戰')]
y = 1.75
for i, (n, d, t, sub) in enumerate(rows):
    box = s.shapes.add_shape(5, Inches(0.9), Inches(y), Inches(11.5), Inches(1.12))
    box.fill.solid(); box.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(0xF3, 0xEE, 0xE1)
    box.line.color.rgb = RGBColor(0xE6, 0xDE, 0xCC); box.shadow.inherit = False
    tb(s, 1.05, y + 0.26, 0.7, 0.6, n, 28, ORANGE, True)
    tb(s, 1.7, y + 0.12, 2.1, 0.5, d, 15, INK_SOFT)
    tb(s, 1.7, y + 0.52, 4.6, 0.5, t, 22, INK, True)
    tb(s, 6.6, y + 0.36, 5.6, 0.5, sub, 17, INK_SOFT)
    y += 1.28
footer(s, 2)

# ── 3 第1節 目標 ───────────────────────────────────────
s = slide(); head(s, '2-1', '第 1 節｜磁鐵能吸引什麼？', '今天的三個任務')
bullets(s, ['① 先猜一猜：哪些東西會被磁鐵吸住？',
            '② 動手實測，把結果記錄在學習單',
            '③ 找出規律：被吸住的東西有什麼共同點？'])
footer(s, 3)

# ── 4 生活中的磁鐵 ─────────────────────────────────────
s = slide(); head(s, '2-1', '磁鐵藏在哪裡？', '這些東西裡面都有磁鐵')
pic_grid(s, ['108', '119', '136', '111'], 0.9, 1.85, 11.5, 3.6, cols=4)
tb(s, 0.9, 5.7, 11.5, 0.9, '想一想：它們各自用磁鐵來做什麼？如果拿掉磁鐵，會發生什麼事？', 19, INK_SOFT)
footer(s, 4)

# ── 5 預測 ────────────────────────────────────────────
s = slide(); head(s, '2-1', '先猜一猜，再動手', '你覺得哪幾個會被吸住？')
pic_grid(s, ['002', '004', '008', '009', '012', '013', '017', '020'], 0.9, 1.8, 11.5, 3.9, cols=4)
tb(s, 0.9, 5.95, 11.5, 0.6, '橡皮擦・長尾夾・鐵釘・硬幣・迴紋針・塑膠尺・鐵湯匙・紙張', 17, INK_SOFT)
footer(s, 5)

# ── 6 公平測量 ────────────────────────────────────────
s = slide(); head(s, '2-1', '怎麼測才公平？', '每個人都用同樣的方法，結果才能比較')
bullets(s, ['用「同一個」磁鐵來測',
            '磁鐵輕輕靠近就好，不要用力壓',
            '每一樣都測三次，結果一樣才算數',
            '被吸住 → 打「✓」；沒被吸住 → 打「✗」'], gap=1.15)
footer(s, 6)

# ── 7 實驗結果 ────────────────────────────────────────
s = slide(); head(s, '2-1', '實驗結果：分成兩堆')
tb(s, 0.9, 1.75, 5.4, 0.5, '✓ 會被吸住', 22, NRED, True)
pic_grid(s, ['004', '008', '012', '017'], 0.9, 2.3, 5.4, 2.4, cols=2)
tb(s, 7.0, 1.75, 5.4, 0.5, '✗ 不會被吸住', 22, SBLUE, True)
pic_grid(s, ['002', '009', '013', '020'], 7.0, 2.3, 5.4, 2.4, cols=2)
tb(s, 0.9, 5.05, 11.5, 0.6, '長尾夾、鐵釘、迴紋針、鐵湯匙　｜　橡皮擦、硬幣、塑膠尺、紙張', 17, INK_SOFT)
tb(s, 0.9, 5.75, 11.5, 0.6, '你的那一組，結果也一樣嗎？', 19, INK, True)
footer(s, 7)

# ── 8 結論 ────────────────────────────────────────────
s = slide(); head(s, '2-1', '結論：磁鐵吸引鐵製品')
tb(s, 1.0, 2.0, 11.0, 1.0, '被磁鐵吸住的物品，都是「鐵」做的。', 32, INK, True)
bullets(s, ['磁鐵吸引的是「鐵製品」，不是所有金屬',
            '銅、鋁、金、銀做出來的東西，磁鐵吸不住'], y=3.3, size=22, gap=1.0)
tb(s, 1.0, 5.4, 11.0, 0.8, '所以：看到金屬，不能直接說「會被吸」——要用磁鐵測過才知道。', 19, ORANGE)
footer(s, 8)

# ── 9 迷思 硬幣 ───────────────────────────────────────
s = slide(); head(s, '2-1', '常見迷思：硬幣是鐵做的嗎？', '摸起來又硬又冰，不代表是鐵')
pic_grid(s, ['009', '016', '068'], 0.9, 1.85, 6.2, 3.4, cols=3)
tb(s, 7.4, 2.0, 5.2, 3.4,
   '臺灣的 1 元、5 元、10 元硬幣，\n主要是銅、鎳、鋅，\n不是鐵。\n\n所以磁鐵吸不住硬幣。', 21, INK, spacing=1.5)
tb(s, 0.9, 5.6, 11.5, 0.6, '小提醒：上課一定要用真的磁鐵試一次，眼見為憑。', 18, INK_SOFT)
footer(s, 9)

# ── 10 第2節 目標 ─────────────────────────────────────
s = slide(); head(s, '2-1・2-2', '第 2 節｜磁力有多大？磁極在哪裡？', '今天的發現')
bullets(s, ['磁鐵的「哪裡」吸力最強？',
            '距離遠一點，還吸得住嗎？',
            '每一塊磁鐵都有 N 極和 S 極'], gap=1.4)
footer(s, 10)

# ── 11 磁鐵的形狀 ─────────────────────────────────────
s = slide(); head(s, '2-1', '磁鐵有很多種形狀')
pic_grid(s, ['025', '026', '027', '088'], 0.9, 1.85, 11.5, 3.5, cols=4)
tb(s, 0.9, 5.5, 11.5, 0.5, '長條形　｜　馬蹄形　｜　圓形　｜　環形', 19, INK_SOFT, align=PP_ALIGN.CENTER)
tb(s, 0.9, 6.1, 11.5, 0.5, '形狀不同，都是磁鐵，都同時有 N 極和 S 極。', 19, INK, True, align=PP_ALIGN.CENTER)
footer(s, 11)

# ── 12 磁力哪裡最強 ───────────────────────────────────
s = slide(); head(s, '2-1', '磁力哪裡最強？', '用迴紋針來比一比')
pic_grid(s, ['095', '096'], 0.9, 1.8, 7.2, 3.2, cols=2)
tb(s, 8.5, 2.0, 4.2, 3.0, '兩端（磁極）\n吸得最多、最牢。\n\n中間幾乎吸不住。', 22, INK, spacing=1.5)
tb(s, 0.9, 5.35, 11.5, 0.6, '結論：磁鐵的磁力集中在兩端，我們把這兩端叫做「磁極」。', 20, ORANGE, True)
footer(s, 12)

# ── 13 鐵粉看磁力 ─────────────────────────────────────
s = slide(); head(s, '2-1', '用鐵粉「看見」磁力', '鐵粉會被吸到磁力最強的地方')
pic_grid(s, ['030', '050'], 0.9, 1.85, 6.6, 3.5, cols=2)
tb(s, 7.8, 2.1, 4.8, 3.2, '鐵粉在磁鐵兩端\n聚得最多、最密。\n\n越往中間，鐵粉越少。', 22, INK, spacing=1.5)
tb(s, 0.9, 5.7, 11.5, 0.6, '如果沒有鐵粉：把迴紋針排在桌上，慢慢把磁鐵靠近，也能看到一樣的現象。', 17, INK_SOFT)
footer(s, 13)

# ── 14 越遠越弱 ───────────────────────────────────────
s = slide(); head(s, '2-1', '距離越遠，磁力越弱')
pic_grid(s, ['107', '120'], 0.9, 1.85, 7.0, 3.3, cols=2)
tb(s, 8.3, 2.2, 4.3, 2.8, '貼著吸：\n一次吸起好幾根\n\n拉遠一點：\n吸到的變少了', 21, INK, spacing=1.45)
tb(s, 0.9, 5.5, 11.5, 0.6, '動手做：把迴紋針排成一排，量量看磁鐵離多遠還吸得到。', 19, ORANGE)
footer(s, 14)

# ── 15 隔著物品 ───────────────────────────────────────
s = slide(); head(s, '2-1', '隔著東西，還吸得住嗎？')
tb(s, 1.0, 1.95, 11.0, 1.0, '磁鐵隔著紙、隔著書、隔著塑膠片，還是能吸引鐵製品。', 27, INK, True, spacing=1.4)
pic_grid(s, ['004', '011', '036'], 0.9, 3.4, 7.2, 2.6, cols=3)
tb(s, 8.6, 3.5, 4.0, 2.2, '生活中的例子：\n冰箱門上的磁鐵\n隔著紙就把便條紙吸住。', 18, INK_SOFT, spacing=1.4)
footer(s, 15)

# ── 16 磁極 N/S ───────────────────────────────────────
s = slide(); head(s, '2-2', '每一塊磁鐵都有兩個磁極', 'N 極和 S 極')
pic_grid(s, ['066', '067', '098'], 0.9, 1.9, 8.0, 3.3, cols=3)
tb(s, 9.3, 2.1, 3.2, 3.0, 'N 極\n\nS 極\n\n斷成兩半，\n還是各有 N、S。', 20, INK, spacing=1.3)
tb(s, 0.9, 5.55, 11.5, 0.6, '把磁鐵吊起來自由轉動，停下來時：N 極指北、S 極指南。', 19, ORANGE)
footer(s, 16)

# ── 17 標磁極 ─────────────────────────────────────────
s = slide(); head(s, '2-2', '一起把磁極標出來', 'N 塗紅色、S 塗藍色')
pic_grid(s, ['100', '101', '102', '104'], 0.9, 1.85, 11.5, 3.4, cols=4)
tb(s, 0.9, 5.5, 11.5, 0.6, '注意：同一種形狀的磁鐵，N、S 的方向可能不一樣，要看標示。', 18, INK_SOFT)
footer(s, 17)

# ── 18 第3節 目標 ─────────────────────────────────────
s = slide(); head(s, '2-2', '第 3 節｜同極相斥、異極相吸', '今天要玩的兩件事')
bullets(s, ['把兩個磁極靠在一起，看看會發生什麼事',
            '做一個會自己轉、指出南北的磁鐵'], gap=1.4)
footer(s, 18)

# ── 19 相斥相吸 ───────────────────────────────────────
s = slide(); head(s, '2-2', '同極相斥、異極相吸')
pic_grid(s, ['091', '092'], 0.9, 1.85, 7.4, 3.3, cols=2)
tb(s, 8.7, 1.95, 4.0, 3.6, 'N－N 互推\nS－S 互推\nN－S 相吸', 22, INK, spacing=1.5)
tb(s, 0.9, 5.45, 11.5, 0.7, '口訣：相同的推開，不同的黏在一起。', 23, ORANGE, True)
footer(s, 19)

# ── 20 相斥小車 ───────────────────────────────────────
s = slide(); head(s, '2-2', '相斥也能做事：磁力小車')
pic_grid(s, ['093', '090'], 0.9, 1.85, 7.0, 3.4, cols=2)
tb(s, 8.2, 2.0, 4.4, 3.4, '把兩台小車的磁鐵\n「同極」相對，\n車子會被推開。\n\n這就是磁浮列車的原理。', 20, INK, spacing=1.45)
footer(s, 20)

# ── 21 指南北 ─────────────────────────────────────────
s = slide(); head(s, '2-2', '磁鐵會指南北')
tb(s, 1.0, 1.95, 7.0, 2.6, '把磁鐵吊起來、\n或放在水面上讓它自由轉動，\n停下來時會指向南北。', 25, INK, spacing=1.45)
pic_grid(s, ['105', '106'], 0.9, 4.4, 7.0, 1.9, cols=2)
tb(s, 8.3, 2.4, 4.2, 3.0, '指南針\n裡面的指針\n就是一塊小磁鐵。', 20, INK_SOFT, spacing=1.5)
footer(s, 21)

# ── 22 生活應用 ───────────────────────────────────────
s = slide(); head(s, '2-3', '第 4 節｜生活中的磁鐵', '找一找：磁鐵藏在哪裡？')
pic_grid(s, ['108', '109', '110', '137', '136', '119'], 0.9, 1.8, 11.5, 3.4, cols=6)
tb(s, 0.9, 5.45, 11.5, 0.6, '門擋・鉛筆盒・櫃子門扣・磁性綁帶・冰箱磁條・黑板', 18, INK_SOFT, align=PP_ALIGN.CENTER)
tb(s, 0.9, 6.05, 11.5, 0.5, '想一想：它用磁鐵解決了什麼問題？', 20, ORANGE, True, align=PP_ALIGN.CENTER)
footer(s, 22)

# ── 23 玩具分類 ───────────────────────────────────────
s = slide(); head(s, '2-3', '哪些玩具有磁鐵？')
tb(s, 0.9, 1.72, 5.4, 0.45, '有磁鐵', 20, NRED, True)
pic_grid(s, ['123', '135', '127', '121'], 0.9, 2.2, 5.4, 2.2, cols=4, gap=0.14)
tb(s, 6.9, 1.72, 5.5, 0.45, '沒有磁鐵', 20, SBLUE, True)
pic_grid(s, ['122', '126', '128', '147'], 6.9, 2.2, 5.5, 2.2, cols=4, gap=0.14)
tb(s, 0.9, 4.75, 11.5, 0.6, '磁鐵釣魚・磁力火車・磁鐵飛鏢・磁鐵跳棋　｜　風車・風箏・陀螺・飛盤', 16, INK_SOFT)
tb(s, 0.9, 5.45, 11.5, 1.0, '特別注意：膠帶黏頭髮是「靜電」，不是磁力。\n磁鐵只吸引鐵製品，不會吸引頭髮、紙屑。', 18, ORANGE)
footer(s, 23)

# ── 24 統整 ───────────────────────────────────────────
s = slide(); head(s, '統整', '磁鐵三件事，你記得了嗎？')
bullets(s, ['① 磁鐵吸引鐵製品（不是所有金屬）',
            '② 兩端磁力最強，有 N 極和 S 極',
            '③ 同極相斥、異極相吸'], y=1.85, size=24, gap=1.15)
tb(s, 0.9, 5.35, 11.5, 0.5, '挑戰：打開教學網頁，做 10 題測驗，再找同學兩人對戰！', 20, ORANGE, True)
footer(s, 24)

prs.save(OUT)
print('已儲存', OUT, '頁數', len(prs.slides.__iter__.__self__._sldIdLst))
